import os, uuid, io, unicodedata
from datetime import date, datetime
from flask import (Flask, render_template, request, redirect, url_for, flash,
                   send_from_directory, send_file, abort)
from flask_sqlalchemy import SQLAlchemy
from werkzeug.utils import secure_filename
from openpyxl import Workbook, load_workbook
from config import ENTIDADES, TIPOS_ARCHIVO, EXTENSIONES

BASE = os.path.dirname(os.path.abspath(__file__))
UPLOADS = os.path.join(BASE, "uploads")
os.makedirs(UPLOADS, exist_ok=True)
os.makedirs(os.path.join(BASE, "data"), exist_ok=True)

app = Flask(__name__)
app.config.update(
    SQLALCHEMY_DATABASE_URI="sqlite:///" + os.path.join(BASE, "data", "residuos.db"),
    MAX_CONTENT_LENGTH=25 * 1024 * 1024,
    SECRET_KEY=os.environ.get("SECRET_KEY", "cambiar-en-produccion"),
)
db = SQLAlchemy(app)

# ---- Modelos generados desde config.py --------------------------------------
MODELOS = {}
for k, e in ENTIDADES.items():
    cols = {"__tablename__": k, "id": db.Column(db.Integer, primary_key=True),
            "creado": db.Column(db.DateTime, default=datetime.utcnow)}
    for n, _l, t in e["campos"]:
        cols[n] = db.Column(db.Date if t == "date" else db.Float if t == "number" else db.Text)
    MODELOS[k] = type(k.capitalize(), (db.Model,), cols)


class Archivo(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    entidad = db.Column(db.String(40), index=True)
    registro_id = db.Column(db.Integer, index=True)
    tipo = db.Column(db.String(60))
    nombre = db.Column(db.String(255))
    guardado = db.Column(db.String(255))
    nota = db.Column(db.Text)
    subido = db.Column(db.DateTime, default=datetime.utcnow)

    @property
    def es_imagen(self):
        return self.guardado.rsplit(".", 1)[-1].lower() in {"jpg", "jpeg", "png", "gif", "webp"}


with app.app_context():
    db.create_all()


# ---- Utilidades -----------------------------------------------------------------
def ent(k):
    if k not in ENTIDADES:
        abort(404)
    return ENTIDADES[k], MODELOS[k]


def convertir(t, v):
    v = (v or "").strip() if isinstance(v, str) or v is None else v
    if v in (None, ""):
        return None
    if t == "date":
        if isinstance(v, datetime): return v.date()
        if isinstance(v, date): return v
        return date.fromisoformat(str(v)[:10])
    if t == "number":
        return float(str(v).replace(",", "."))
    return str(v).strip()


def norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode()
    return " ".join(s.lower().split())


def fmt(v):
    if isinstance(v, date): return v.strftime("%d/%m/%Y")
    if isinstance(v, float): return str(int(v)) if v == int(v) else f"{v:g}"
    return "" if v is None else v


app.jinja_env.filters["fmt"] = fmt
app.jinja_env.globals.update(ENTIDADES=ENTIDADES, TIPOS_ARCHIVO=TIPOS_ARCHIVO)


def pivot(M, fila, col):
    datos, cols = {}, set()
    for r in M.query.all():
        f, c = getattr(r, fila) or "(vacío)", getattr(r, col) or "(vacío)"
        datos.setdefault(f, {}); datos[f][c] = datos[f].get(c, 0) + (r.cantidad or 0)
        cols.add(c)
    cols = sorted(cols)
    return {"cols": cols, "filas": {f: [d.get(c, 0) for c in cols] for f, d in sorted(datos.items())}}


# ---- Rutas ----------------------------------------------------------------------
@app.route("/")
def dashboard():
    M = MODELOS["inventario"]
    resumen = {k: m.query.count() for k, m in MODELOS.items()}
    return render_template("dashboard.html", resumen=resumen, n_archivos=Archivo.query.count(),
                           p1=pivot(M, "tipo_residuo", "tipo_contenedor"),
                           p2=pivot(M, "tipo_contenedor", "estado_contenedor"))


@app.route("/<k>/")
def listado(k):
    e, M = ent(k)
    q = request.args.get("q", "").strip()
    query = M.query
    if q:
        txt = [getattr(M, n).ilike(f"%{q}%") for n, _l, t in e["campos"] if t not in ("date", "number")]
        query = query.filter(db.or_(*txt))
    filas = query.order_by(M.id.desc()).limit(500).all()
    adj = dict(db.session.query(Archivo.registro_id, db.func.count()).filter_by(entidad=k)
               .group_by(Archivo.registro_id).all())
    return render_template("list.html", k=k, e=e, filas=filas, q=q, adj=adj)


@app.route("/<k>/nuevo", methods=["GET", "POST"])
@app.route("/<k>/<int:rid>/editar", methods=["GET", "POST"])
def formulario(k, rid=None):
    e, M = ent(k)
    obj = db.get_or_404(M, rid) if rid else None
    if request.method == "POST":
        obj = obj or M()
        try:
            for n, _l, t in e["campos"]:
                setattr(obj, n, convertir(t, request.form.get(n)))
        except ValueError:
            flash("Hay un valor numérico o de fecha inválido.", "err")
            return render_template("form.html", k=k, e=e, obj=obj)
        db.session.add(obj); db.session.commit()
        flash("Registro guardado.", "ok")
        return redirect(url_for("detalle", k=k, rid=obj.id))
    return render_template("form.html", k=k, e=e, obj=obj)


@app.route("/<k>/<int:rid>")
def detalle(k, rid):
    e, M = ent(k)
    obj = db.get_or_404(M, rid)
    archivos = Archivo.query.filter_by(entidad=k, registro_id=rid).order_by(Archivo.id.desc()).all()
    return render_template("detail.html", k=k, e=e, obj=obj, archivos=archivos)


@app.post("/<k>/<int:rid>/eliminar")
def eliminar(k, rid):
    _e, M = ent(k)
    obj = db.get_or_404(M, rid)
    for a in Archivo.query.filter_by(entidad=k, registro_id=rid).all():
        borrar_archivo(a)
    db.session.delete(obj); db.session.commit()
    flash("Registro eliminado.", "ok")
    return redirect(url_for("listado", k=k))


def borrar_archivo(a):
    try: os.remove(os.path.join(UPLOADS, a.guardado))
    except OSError: pass
    db.session.delete(a)


@app.post("/<k>/<int:rid>/subir")
def subir(k, rid):
    _e, M = ent(k)
    db.get_or_404(M, rid)
    n = 0
    for f in request.files.getlist("archivos"):
        nombre = secure_filename(f.filename or "")
        ext = nombre.rsplit(".", 1)[-1].lower() if "." in nombre else ""
        if not nombre or ext not in EXTENSIONES:
            flash(f"Archivo no permitido: {f.filename}", "err"); continue
        guardado = f"{uuid.uuid4().hex}.{ext}"
        f.save(os.path.join(UPLOADS, guardado))
        db.session.add(Archivo(entidad=k, registro_id=rid, nombre=nombre, guardado=guardado,
                               tipo=request.form.get("tipo") or "Otro", nota=request.form.get("nota")))
        n += 1
    db.session.commit()
    if n: flash(f"{n} archivo(s) subido(s).", "ok")
    return redirect(url_for("detalle", k=k, rid=rid))


@app.route("/archivo/<int:aid>")
def ver_archivo(aid):
    a = db.get_or_404(Archivo, aid)
    return send_from_directory(UPLOADS, a.guardado, download_name=a.nombre,
                               as_attachment=not a.es_imagen and not a.guardado.endswith(".pdf"))


@app.post("/archivo/<int:aid>/eliminar")
def quitar_archivo(aid):
    a = db.get_or_404(Archivo, aid)
    k, rid = a.entidad, a.registro_id
    borrar_archivo(a); db.session.commit()
    return redirect(url_for("detalle", k=k, rid=rid))


@app.route("/<k>/exportar")
def exportar(k):
    e, M = ent(k)
    wb = Workbook(); ws = wb.active; ws.title = e["titulo"][:31]
    ws.append([l for _n, l, _t in e["campos"]])
    for r in M.query.order_by(M.id).all():
        ws.append([getattr(r, n) for n, _l, _t in e["campos"]])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return send_file(buf, download_name=f"{k}_{date.today()}.xlsx", as_attachment=True)


@app.route("/<k>/importar", methods=["GET", "POST"])
def importar(k):
    e, M = ent(k)
    if request.method == "POST":
        f = request.files.get("excel")
        if not f or not f.filename.lower().endswith(".xlsx"):
            flash("Sube un archivo .xlsx", "err")
            return render_template("import.html", k=k, e=e)
        ws = load_workbook(f, read_only=True, data_only=True).worksheets[0]
        mapa = {norm(l): (n, t) for n, l, t in e["campos"]}
        cab, total, errores = None, 0, 0
        for i, fila in enumerate(ws.iter_rows(values_only=True)):
            if cab is None:
                m = [mapa.get(norm(c)) for c in fila]
                if sum(1 for x in m if x) >= 3: cab = m
                elif i > 15: break
                continue
            if not any(c not in (None, "") for c in fila) or str(fila[0] or "").upper().startswith("TOTAL"):
                continue
            obj = M()
            try:
                for c, v in zip(cab, fila):
                    if c: setattr(obj, c[0], convertir(c[1], v))
                db.session.add(obj); total += 1
            except (ValueError, TypeError):
                errores += 1
        db.session.commit()
        if cab is None: flash("No se encontró una fila de encabezados reconocible.", "err")
        else: flash(f"Importados {total} registros ({errores} filas omitidas por formato).", "ok")
        return redirect(url_for("listado", k=k))
    return render_template("import.html", k=k, e=e)


if __name__ == "__main__":
    app.run(debug=os.environ.get("FLASK_DEBUG") == "1", host="0.0.0.0", port=5000)
